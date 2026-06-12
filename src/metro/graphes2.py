""" Trace la seconde feuille de graphes"""

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference

from metro.forme import zone_graph, GRAS


def bar_horiz_graph(valeurs:Reference, categories:Reference) -> BarChart:
    """
    Trace un diagramme en barres horizontales
    :param valeurs: valeurs de l'axe horizontal
    :param categories: étiquettes de l'axe vertical
    """
    retour = BarChart()
    retour.type = "bar"
    retour.add_data(valeurs, titles_from_data=True)
    retour.set_categories(categories)
    retour.legend = None
    retour.width = 17
    retour.height = 12
    return retour


def graphes2(classeur: Workbook) -> None:
    """
    Ajoute la seconde feuille de graphes ('Graphes2') au classeur
    :param classeur: classeur contenant les feuilles de données telles que générées par 'feuilles.py',
                    contenant donc la feuille 'stations'.
                    C'est également le classeur où sera ajoutée la feuille 'Graphes2'.
    """
    feuille = classeur.create_sheet("Graphes2")

    # Mise en page
    zone_graph(feuille, "Les 10 stations les plus et les moins fréquentées par les "
                            "détenteurs de certains titres de transport")
    feuille["B2"].value = "Titres de transport \"sociaux\" :"
    feuille["B2"].font = GRAS
    feuille["M2"].value = "Titres de transport de courte durée :"
    feuille["M2"].font = GRAS

    feuille.add_chart(bar_horiz_graph(Reference(classeur["stations"], min_col=23, max_col=24, min_row=1, max_row=21),
                                      Reference(classeur["stations"], min_col=13, min_row=2, max_row=21)), "B3")
    feuille.add_chart(bar_horiz_graph(Reference(classeur["stations"], min_col=37, max_col=38, min_row=1, max_row=21),
                                      Reference(classeur["stations"], min_col=26, min_row=2, max_row=21)), "M3")
