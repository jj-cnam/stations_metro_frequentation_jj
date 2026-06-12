"""Définit les constantes de mise en forme des cellules et la procédure de mise en page des feuilles de graphes"""

from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.styles import Alignment, Font, PatternFill


# Constantes de mise en forme des  cellules
VERT_JADE = PatternFill(fill_type="solid", start_color="00C4B3")
BLEU =      PatternFill(fill_type="solid", start_color="0A0082")
CENTRE =    Alignment(horizontal="center", vertical="center")
DROITE =    Alignment(horizontal="right")
TITRE1 =    Font(size=24, bold=True, color="FFFFFF")
TEXTE =     Font(size=16, color="FFFFFF")
GRAS =      Font(size=16, bold=True, color="FFFFFF")


def zone_graph(feuille: Worksheet, titre:str, ligne1:int = 1) -> None:
    """
    Procédure qui met en forme une zone de graphes :

    - cartouche vert jade
    - titre principal mis en valeur
    - fond bleu

    :param feuille: Feuille sur laquelle appliquer la mise en forme
    :param titre: Titre principal
    :param ligne1: Numéro de la première ligne
    """
    cartouche = "A{}:W{}".format(ligne1, ligne1)
    titre_cell = "A{}".format(ligne1)
    plage = "A{}:W{}".format(ligne1+1, ligne1+31)
    for lignes in feuille[cartouche]:
        for cellule in lignes:
            cellule.fill = VERT_JADE
    feuille.merge_cells(cartouche)
    feuille[titre_cell].value = titre
    feuille[titre_cell].alignment = CENTRE
    feuille[titre_cell].font = TITRE1
    for lignes in feuille[plage]:
        for cellule in lignes:
            cellule.fill = BLEU
