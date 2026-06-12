"""Trace la première feuille de graphes"""

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.chart import BarChart, Reference
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from metro.forme import zone_graph, TEXTE, DROITE


def cellule_selection(feuille: Worksheet, cellule: str, nomdef:str, plage:str) -> None:
    """
    Procédure qui créée une cellule avec un nom défini, qui présente une liste déroulante des choix autorisés.
    Cette cellule sera celle sélectionée par défaut à l'ouverture de la feuille.
    :param feuille: feuille contenant la cellule
    :param cellule: cellule à définir, en **coordonnées absolues ($c$l)**
    :param nomdef: nom défini
    :param plage: plage des valeurs autorisées
    """
    feuille.parent.defined_names.add(DefinedName(nomdef, attr_text=feuille.title + "!" + cellule))
    liste_val = DataValidation(type="list", formula1=plage, allowBlank=False, showDropDown=False)
    feuille.add_data_validation(liste_val)
    liste_val.add(feuille[cellule])
    feuille.sheet_view.selection[0].activeCell = cellule
    feuille.sheet_view.selection[0].sqref = cellule


def bar_vertic_graph(valeurs:Reference, categories:Reference,
                     pourcent:bool = False, titreX:str|None = None) -> BarChart:
    """
    Trace un diagramme en barres verticales évec empilement
    :param valeurs: valeurs de l'axe vertical
    :param categories: étiquettes de l'axe horizontal
    :param pourcent: True pour avoir un empilement en pourcentage
    :param titreX: Titre de l'axe horizontal
    """
    retour = BarChart(overlap=100)
    if pourcent:
        retour.grouping = "percentStacked"
    else:
        retour.grouping = "stacked"
    retour.add_data(valeurs, titles_from_data=True)
    retour.set_categories(categories)
    # Titre dynamique : pas possible en OpenPyxl
    # Calcul auto de la taille : pas possible en OpenPyxl
    retour.width = 35
    retour.height= 12
    retour.x_axis.delete = False
    if titreX: 
        retour.x_axis.title = titreX
    retour.y_axis.delete = False
    retour.legend.position = "t"
    return retour


def graphes1(classeur: Workbook) -> None:
    """
    Ajoute la première feuille de graphes ('Graphes1') au classeur
    :param classeur: classeur contenant les feuilles de données telles que générées par 'feuilles.py',
                    contenant donc les feuilles 'semaines' et 'echantillon'.
                    C'est également le classeur où sera ajoutée la feuille 'Graphes1'.
    """
    feuille = classeur.create_sheet("Graphes1")

    ## Mise en page
    zone_graph(feuille, "Nombre de validations sur le réseau métro")
    feuille["F2"].value = "Sélectionner la station :"
    feuille["F2"].font = TEXTE
    feuille["F2"].alignment = DROITE
    feuille["G2"].font = TEXTE     # cellule de sélection ('station')
    feuille.merge_cells("G2:I2")
    zone_graph(feuille, "Répartition des types de titres de transport selon les stations", ligne1=28)
    feuille["F29"].value = "Échantillon aléatoire de 50 stations. Recalculer le classeur [F9] pour renouveler l'échantillon."
    feuille["F29"].font = TEXTE

    # Cellule de sélection de la station à afficher :
    ## La cellule G2 doit avoir le nom défini 'station', qui est utilisé dans les formules de calcul des fréquentations
    ## journalières et hebdomadaires par station dans 'feuilles.py'
    ## Elle doit ne permettre d'entrer qu'un nom de station, donc compris dans celles présentes dans la
    ## feuille 'modalites', ou la valeur "toutes" :
    cellule_selection(feuille, "$G$2", "station", "=modalites!A1:A319")
    ## La valeur "toutes" sera sa valeur par défaut :
    feuille["$G$2"].value = "toutes"

    feuille.add_chart(bar_vertic_graph(Reference(classeur["semaines"], min_col=2, max_col=8, min_row=1, max_row=54),
                                       Reference(classeur["semaines"], min_col=1, min_row=2, max_row=54),
                                       titreX="Numéro de semaine"), "B3")
    feuille.add_chart(bar_vertic_graph(Reference(classeur["echantillon"], min_col=2, max_col=8, min_row=1, max_row=51),
                                       Reference(classeur["echantillon"], min_col=1, min_row=2, max_row=51),
                                       pourcent=True), "B30")
