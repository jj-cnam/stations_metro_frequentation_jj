"""Calcule les feuilles du classeur"""

import datetime
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.utils import get_column_letter


def unique(ensemble: tuple | list, suppr_entete=True) -> list:     # type de plage à mettre
    """
    Retourne la liste triée des éléments uniques d'un ensemble
    :param ensemble:
    :param suppr_entete: True pour supprimer la première valeur de l'ensemble
    """
    ensemble = list(ensemble)
    if suppr_entete: 
        ensemble.pop(0)
    uniques = set()
    for element in ensemble:
        uniques.add(element.value)
    uniques = list(uniques)
    uniques.sort()
    return uniques


def topflop(feuille: Worksheet, debut:int) -> None:
    """
    Procédure qui renseigne les colonnes 'top' et 'flop' de la feuille 'stations' (voir documentation)
    et les met au format pourcentage
    :param feuille: feuille 'stations'
    :param debut: numéro de la première colonne ('top')
    """
    feuille.cell(1, debut).value = "top"
    for j in range(2, 12):
        feuille.cell(j, debut).value = "=" + get_column_letter(debut-1) + str(j)
        feuille.cell(j, debut).number_format = "0.00%"
    feuille.cell(1, debut+1).value = "flop"
    for j in range(12, 22):
        feuille.cell(j, debut+1).value =  "=" + get_column_letter(debut-1) + str(j)
        feuille.cell(j, debut+1).number_format = "0.00%"


def feuilles(classeur: Workbook):
    # 1. feuille 'DATA' : ajout d'une colonne avec le numéro de la semaine
    data = classeur['DATA']
    data["E1"].value = "semaine"
    for ligne in range(2, data.max_row+1):  # enumerate ?
        data.cell(row=ligne, column=5).value = "=WEEKNUM($A{l})".format(l=ligne)

    # 2. feuille 'modalites'
    modalites = classeur.create_sheet("modalites")
    ## stations
    liste_stations = unique(data["B"])
    for i, valeur in enumerate(liste_stations, start=1):
        modalites.cell(row=i, column=1, value=valeur)
    # la valeur/modalité "toutes" sera ajoutée pour permettre les calculs portant sur l'ensemble des stations
    modalites.append(("toutes",))

    ## titres de transport
    liste_titres = unique(data["C"])
    for i, valeur in enumerate(liste_titres, start=1):
        modalites.cell(row=i, column=2, value=valeur)

    # 3. feuille 'jours'
    jours = classeur.create_sheet("jours")
    ## entêtes
    jours.append(["jour",] + liste_titres)
    ## valeurs : jour + les comptes des différents titres de transport
    ## =SOMME.SI.ENS(DATA!$D:$D;DATA!$A:$A;$A{l};DATA!$B:$B;SI(station="toutes";"*";station);DATA!$C:$C;B$1)
    ## la cellule 'station' (nom défini) est définie dans 'graphes1.py'
    ## avec {l}=numéro de la ligne et {c}=lettre de la colonne (B à H)
    formule_jours = "=SUMIFS(DATA!$D:$D,DATA!$A:$A,$A{l},DATA!$B:$B,IF(station=\"toutes\",\"*\",station),DATA!$C:$C,{c}$1)"
    for i in range(0, 365):
        ligne:list = [datetime.date(2025, 1, 1) + datetime.timedelta(i)]
        for colonne in ("B", "C", "D", "E", "F", "G", "H"):
            ligne.append(formule_jours.format(l=i+2, c=colonne))
        jours.append(ligne)

    # 4. feuille 'semaines'
    semaines = classeur.create_sheet("semaines")
    ## entêtes
    semaines.append(["semaine",] + liste_titres)
    ## valeurs : numéro de semaine + les comptes des différents titres de transport
    ## formule_semaines = formule_jours en remplaçant DATA!$A:$A,$A{l} par DATA!$E:$E,$A{l}
    formule_semaines = \
        "=SUMIFS(DATA!$D:$D,DATA!$E:$E,$A{l},DATA!$B:$B,IF(station=\"toutes\",\"*\",station),DATA!$C:$C,{c}$1)"
    for i in range(1, 54):
        ligne:list = [i]
        for colonne in ("B", "C", "D", "E", "F", "G", "H"):
            ligne.append(formule_semaines.format(l=i+1, c=colonne))
        semaines.append(ligne)

    # 5. feuille 'stations'
    stations = classeur.create_sheet("stations")
    ## entêtes
    stations.append(["station",] + liste_titres + ["aleatoire", "prop_social", "prop_court"])
    ## valeurs : nom de station + les comptes des différents titres de transport +...
    ## =SOMME.SI.ENS(DATA!$D:$D;DATA!$B:$B;$A{l};DATA!$C:$C;{c}$1)
    formule_stations = "=SUMIFS(DATA!$D:$D,DATA!$B:$B,$A{l},DATA!$C:$C,{c}$1)"
    for i in range(len(liste_stations)):
        ligne = [liste_stations[i],]
        for colonne in ("B", "C", "D", "E", "F", "G", "H"):
            ligne.append(formule_stations.format(l=i+2, c=colonne))
        ## ...+ valeur aléatoire qui servira pour l'échantillonnage dans l'onglet suivant +...
        ligne.append("=RAND()")
        ## ...+ proportion de titres de transport 'sociaux' ("Amethyste" et "Contrat Solidarité Transport") et courts
        ligne.append("=(B{l}+D{l})/SUM(B{l}:H{l})".format(l=i+2))
        ligne.append("=(F{l}+D{l})/SUM(B{l}:H{l})".format(l=i+2))
        stations.append(ligne)

    ## chiffres pour les 'titres de transport sociaux'
    ## =PRENDRE(TRIER($A$2:$J$319;10;-1;);10)
    stations["M2"] = ArrayFormula("M2:V11",
                                  "=_xlfn.TAKE(_xlfn.SORT($A$2:$J$319,10,-1,),10)")
    stations["M12"] = ArrayFormula("M12:V21",
                                  "=_xlfn.TAKE(_xlfn.SORT($A$2:$J$319,10,-1,),-10)")
    topflop(stations, 23)

    ## chiffres pour les 'titres de transport courts'
    stations["Z2"] = ArrayFormula("Z2:AJ11",
                                  "=_xlfn.TAKE(_xlfn.SORT($A$2:$K$319,11,-1,),10)")
    stations["Z12"] = ArrayFormula("Z12:AJ21",
                                  "=_xlfn.TAKE(_xlfn.SORT($A$2:$K$319,11,-1,),-10)")
    topflop(stations, 37)

    # 6. feuille 'echantillon' : tirage aléatoire dynamique de 50 lignes de la feuille 'stations'
    # en se basant sur les 50 premières valeurs de la colonne 'aleatoire'
    echantillon = classeur.create_sheet("echantillon")
    ## entêtes
    echantillon.append(["station",] + liste_titres)

    ## =PRENDRE(TRIERPAR(stations!A:H;stations!I:I);50)
    classeur["echantillon"]["A2"] = ArrayFormula(
        'A2:H51', "=_xlfn.TAKE(_xlfn.SORTBY(stations!A:H,stations!I:I),50)")
